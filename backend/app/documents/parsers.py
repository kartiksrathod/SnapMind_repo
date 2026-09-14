from dataclasses import dataclass
from pathlib import Path


@dataclass
class ParsedSegment:
    text: str
    page: int | None = None
    section: str | None = None


@dataclass
class ParsedDocument:
    text: str
    page_count: int | None = None
    segments: list[ParsedSegment] | None = None


class DocumentParseError(Exception):
    pass


class DocumentParser:
    extensions: tuple[str, ...] = ()

    def parse(self, path: Path) -> ParsedDocument:
        raise NotImplementedError


class TextParser(DocumentParser):
    extensions = ('.txt', '.md', '.markdown')

    def parse(self, path: Path) -> ParsedDocument:
        try:
            text = path.read_text(encoding='utf-8-sig')
        except UnicodeDecodeError as exc:
            raise DocumentParseError('The text file is not valid UTF-8.') from exc
        if not text.strip():
            raise DocumentParseError('No usable text was found in this document.')
        return ParsedDocument(text=text, segments=[ParsedSegment(text=text)])


class PdfParser(DocumentParser):
    extensions = ('.pdf',)

    def parse(self, path: Path) -> ParsedDocument:
        try:
            from pypdf import PdfReader
            reader = PdfReader(str(path))
            segments = [ParsedSegment(text=page.extract_text() or '', page=page_number) for page_number, page in enumerate(reader.pages, 1)]
            text = '\n'.join(segment.text for segment in segments).strip()
        except Exception as exc:
            raise DocumentParseError('The PDF could not be read. It may be corrupted or image-only.') from exc
        if not text:
            raise DocumentParseError('No usable text was found in this PDF.')
        return ParsedDocument(text=text, page_count=len(reader.pages), segments=segments)


class DocxParser(DocumentParser):
    extensions = ('.docx',)

    def parse(self, path: Path) -> ParsedDocument:
        try:
            from docx import Document
            document = Document(str(path))
            text = '\n'.join(paragraph.text for paragraph in document.paragraphs).strip()
        except Exception as exc:
            raise DocumentParseError('The DOCX could not be read. It may be corrupted.') from exc
        if not text:
            raise DocumentParseError('No usable text was found in this DOCX.')
        return ParsedDocument(text=text, segments=[ParsedSegment(text=text)])


PARSERS = {extension: parser for parser in (TextParser(), PdfParser(), DocxParser()) for extension in parser.extensions}


def parser_for(extension: str) -> DocumentParser:
    try:
        return PARSERS[extension.lower()]
    except KeyError as exc:
        raise DocumentParseError(f'Unsupported document format: {extension or "unknown"}.') from exc
